from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "Informe_avance_validacion_semanas_4_y_5.docx"
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.72); section.bottom_margin = Inches(.7)
section.left_margin = Inches(.8); section.right_margin = Inches(.8)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
for name, size in [("Title", 20), ("Heading 1", 13), ("Heading 2", 11)]:
    styles[name].font.name = "Aptos Display"; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0,0,0)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]; cell.text = h; shade(cell, "1F4E78"); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.font.bold = True
    for row_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_i % 2: shade(cells[i], "EAF2F8")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t
def bullet(text):
    p = doc.add_paragraph(style="List Bullet"); p.add_run(text); return p
def code(text):
    p = doc.add_paragraph(); r=p.add_run(text); r.font.name="Consolas"; r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas"); r.font.size=Pt(8.5); return p

title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run("Informe de avance de validación experimental")
subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER; subtitle.add_run("Semanas 4 y 5  Prototipo offline first de vales").italic = True
p=doc.add_paragraph(); p.add_run("Conclusión. ").bold=True; p.add_run("El laboratorio, el dataset y una primera ejecución integral de EP01 a EP04 quedaron operativos. La evidencia confirma registro offline, sincronización concurrente, trazabilidad y recuperación idempotente. El resultado es una línea base funcional; la Semana 6 debe repetir las corridas para producir evidencia estadística.")

doc.add_heading("Alcance logrado", level=1)
table(["Semana", "Entregable", "Resultado"], [
    ["4", "Laboratorio e instrumentación", "Dataset reproducible de 1,000 operaciones, correlación por lote, logs JSONL, reintentos y fallo inyectable."],
    ["5", "Primera corrida experimental", "EP01 a EP04 ejecutados contra uvicorn real y base central aislada."],
], [0.7, 2.2, 3.7])

doc.add_heading("Dataset y comportamiento simulado", level=1)
doc.add_paragraph("El generador distribuye 200 operaciones en cada uno de los cinco nodos CLI 1 a CLI 5. Incluye 850 agregados exclusivos y 150 operaciones sobre 50 agregados compartidos, con cambios de cantidad, material y cancelación. Esta reserva de solapamientos ejercita la unión CRDT en condiciones concurrentes.")
code("tests/lab/generar_dataset.py  |  outputs/plan/dataset.json")
bullet("EP01 registra los 1,000 vales en SQLite local sin sincronizar inicialmente.")
bullet("EP02 y EP03 envían cinco lotes concurrentes al mismo servidor uvicorn y guardan las marcas de envío y confirmación.")
bullet("EP04 corta la respuesta después de que el servidor confirma el lote; el reintento devuelve duplicate y comprueba idempotencia por event_id.")
code("tests/lab/execute_plan.py  |  tests/lab/fault_injector.py")

doc.add_heading("Resultados de la primera corrida", level=1)
table(["Verificación", "Resultado", "Interpretación"], [
    ["EP01", "1,000 de 1,000", "Registros offline confirmados localmente."],
    ["EP02", "5 lotes", "Sincronización concurrente completada."],
    ["EP04", "1 fallo y recuperación", "Reintento detectó duplicate sin duplicar el evento central."],
    ["Eventos centrales", "1,001", "1,000 del dataset más 1 evento de recuperación."],
    ["IDs duplicados", "0", "Idempotencia preservada."],
    ["Movimiento auditoría", "1,001 a 1,001", "Trazabilidad directa sin auditorías huérfanas."],
], [1.35, 1.25, 4.0])
doc.add_paragraph("La evidencia cuantitativa se consolidó en outputs/plan/verification.json. Las latencias individuales por lote se exportaron en outputs/plan/lotes.csv para el procesamiento posterior en R Studio o Jamovi.")

doc.add_heading("Puntos críticos del código", level=1)
bullet("La cantidad se representa como PN Counter por nodo y los campos descriptivos se resuelven mediante LWW por campo. La unión se comparte entre cliente y servidor, evitando una regla distinta por réplica.")
code("shared/crdt.py  |  backend/app/services/sync_service.py  |  client/app/sync/conflict_resolver.py")
bullet("El servidor serializa la aplicación de eventos por aggregate_id y valida el payload antes de persistir. El lock es adecuado para SQLite en un proceso; una instalación multi proceso requiere PostgreSQL y bloqueo transaccional.")
bullet("Cada lote incorpora lote_id, enviado_en y confirmado_en. Los JSONL del cliente y servidor son la fuente experimental; los logs técnicos de uvicorn no contienen indicadores de tesis.")
code("shared/schemas.py  |  client/app/sync/sync_engine.py  |  outputs/plan/client-sync.jsonl  |  outputs/plan/server-events.jsonl")
bullet("Eventos y movimientos son append only. Auditoria.movimiento_id enlaza explícitamente cada movimiento con su registro de auditoría.")
code("backend/app/models/evento_sync.py  |  backend/app/models/movimiento_inventario.py  |  backend/app/models/auditoria.py")

doc.add_heading("Estructura de datos y sincronización", level=1)
table(["Base", "Tablas", "Función"], [
    ["Local por nodo", "vales, eventos, nodo_campo", "Estado local, cola pendiente y estado de conexión."],
    ["Central", "vales, eventos_sync, movimientos_inventario, auditoria, nodos_campo", "Estado consolidado, historial, inventario, trazabilidad y disponibilidad."],
], [1.2, 2.9, 2.5])
doc.add_paragraph("La conexión se realiza mediante /api/v1/sync/bulk. El cliente envía eventos pendientes; el servidor los registra en eventos_sync, fusiona el estado del vale, crea movimiento y auditoría, y devuelve eventos remotos para que cada base local converja.")

doc.add_heading("Siguiente paso Semana 6", level=1)
doc.add_paragraph("Se deben repetir EP01 a EP04 en varias corridas independientes, conservando una carpeta de evidencia por corrida. Con esas repeticiones se podrán calcular medidas descriptivas y contrastes estadísticos durante la Semana 7. También conviene ampliar EP04 con cortes antes de llegar al servidor y durante la transferencia de lotes de distinto tamaño.")

footer = section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("Informe de avance  Semanas 4 y 5")
doc.save(OUT)
